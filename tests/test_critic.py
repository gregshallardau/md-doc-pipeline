"""CriticMarkup review marks: resolve, render, lint and round-trip through Word."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from click.testing import CliRunner

from md_doc import critic
from md_doc.cli import main

SAMPLE = """---
title: T
note: "{++frontmatter stays++}"
---
# {++New ++}title

Hello {--old --}{~~colour~>color~~} {==world==}{>>check *this*<<} end.

- {++added item++}
- keep
- {--gone item--}

before

{++
para one

para two
++}

after

`{++code++}`

```
{--fenced--}
```
"""


def test_accept_applies_every_change():
    out = critic.apply(SAMPLE, "accept")
    assert "# New title" in out
    assert "Hello color world end." in out
    assert "- added item\n- keep\n" in out and "gone item" not in out
    assert "para one\n\npara two" in out
    assert "check" not in out


def test_reject_reverts_every_change():
    out = critic.apply(SAMPLE, "reject")
    assert "# title" in out
    assert "Hello old colour world end." in out
    assert "- keep\n- gone item" in out and "added item" not in out
    assert "para one" not in out


def test_code_and_frontmatter_are_left_alone():
    for mode in critic.MODES:
        out = critic.apply(SAMPLE, mode)
        assert 'note: "{++frontmatter stays++}"' in out
        assert "`{++code++}`" in out
        assert "{--fenced--}" in out


def test_removed_marks_leave_no_stray_lines():
    reject = critic.apply("a\n\n{++\nnew para\n++}\n\nb\n- {++x++}\n", "reject")
    assert reject == "a\n\nb\n"
    accept = critic.apply("a\n\n{++\nnew para\n++}\n\nb\n", "accept")
    assert accept == "a\n\nnew para\n\nb\n"
    assert critic.apply("## {--Old heading--}\n\ntext\n", "accept").strip() == "text"


def test_show_renders_inline_html():
    out = critic.apply(SAMPLE, "show")
    assert '# <ins class="critic-ins">New</ins> title' in out
    assert '<del class="critic-del">old</del> <del class="critic-del">colour</del>' in out
    assert '<ins class="critic-ins">color</ins>' in out
    assert '<mark class="critic-hl critic-c1">world</mark>' in out  # paired with its comment
    assert (
        '<span class="critic-comment" data-comment="check *this*" data-c="1" data-mark="0"></span>'
        in out
    )
    assert '- <ins class="critic-ins">added item</ins>' in out
    assert '<ins class="critic-ins">para one</ins>\n\n<ins class="critic-ins">para two</ins>' in out


def test_show_wraps_table_cells_not_the_row():
    out = critic.apply("| a | b |\n|---|---|\n| {++x++} | y |\n", "show")
    assert '| <ins class="critic-ins">x</ins> | y |' in out
    assert "|---|---|" in out


def test_comment_text_is_escaped():
    out = critic.apply('{>>say "hi" <b>now</b><<}', "show")
    assert 'data-comment="say &quot;hi&quot; &lt;b&gt;now&lt;/b&gt;"' in out


def test_slides_cannot_show_review_marks():
    assert critic.apply("{++a++}{--b--}", "show", fmt="pptx") == "a"


def test_modes_and_defaults():
    assert critic.review_mode({}) == "accept"
    assert critic.review_mode({"review_markup": "SHOW"}) == "show"
    assert critic.review_mode({"review_markup": "bogus"}) == "accept"
    assert critic.review_author({}) == "Reviewer"
    assert critic.review_author({"review_author": " Ada L "}) == "Ada L"
    assert critic.apply("no marks here {x}", "show") == "no marks here {x}"


def test_stamp_honours_source_date_epoch(monkeypatch):
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "86400")
    assert critic.review_stamp() == "1970-01-02T00:00:00Z"


def test_problems_report_only_broken_marks():
    found = critic.problems("ok {++a++}\nbad {++ b\n\n{~~x~~}\n`{--code`\n")
    assert [line for line, _ in found] == [2, 4]
    assert "never closed" in found[0][1] and "~>" in found[1][1]
    assert critic.problems("{++a++} {--b--} {~~c~>d~~} {==e==}{>>f<<}") == []


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------

DOC = """---
title: Review
author: Ada
review_author: Greg Shallard
---
# Review

Hello {--old --}{~~colour~>color~~} {==world==}{>>Check this<<} end.

- {++New bullet++}

| A | B |
|---|---|
| {++x++} | {--y--} |
"""


def build_docx(tmp_path: Path, mode: str | None) -> Path:
    from md_doc.builders import docx

    (tmp_path / ".git").mkdir(exist_ok=True)
    config = {"title": "Review", "review_author": "Greg Shallard"}
    if mode:
        config["review_markup"] = mode
    out = tmp_path / f"{mode}.docx"
    docx.build(DOC, config, out, repo_root=tmp_path, doc_path=tmp_path / "d.md")
    return out


def xml_of(path: Path, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode()


def test_word_show_builds_tracked_changes_and_comments(tmp_path):
    from docx import Document

    out = build_docx(tmp_path, "show")
    xml = xml_of(out)
    assert len(re.findall(r"<w:ins ", xml)) == 2  # color, New bullet
    assert len(re.findall(r"<w:del ", xml)) == 2  # old, colour
    assert len(re.findall(r"<w:delText", xml)) == 2
    assert 'w:author="Greg Shallard"' in xml
    assert 'w:highlight w:val="yellow"' in xml
    comments = [(c.author, c.text) for c in Document(str(out)).comments]
    assert comments == [("Greg Shallard", "Check this")]
    assert "commentRangeStart" in xml


def test_word_table_cells_show_marks_as_formatting(tmp_path):
    xml = xml_of(build_docx(tmp_path, "show"))
    assert "<w:strike/>" in xml and "<w:u " in xml


@pytest.mark.parametrize("mode", [None, "accept", "reject"])
def test_word_clean_modes_have_no_revisions(tmp_path, mode):
    out = build_docx(tmp_path, mode)
    xml = xml_of(out)
    assert "<w:ins " not in xml and "<w:del " not in xml
    assert "word/comments.xml" not in zipfile.ZipFile(out).namelist()
    from docx import Document

    text = " ".join(p.text for p in Document(str(out)).paragraphs)
    assert ("Hello color" in text) == (mode != "reject")
    assert ("Hello old colour" in text) == (mode == "reject")


def all_texts(box) -> list[str]:
    """Every text run under *box*, including absolutely positioned margin notes."""
    found = [box.text] if type(box).__name__ == "TextBox" else []
    for child in getattr(box, "children", []) or []:
        found += all_texts(child)
    return found


def test_pdf_show_renders_marks_comment_notes_and_a_legend(tmp_path, monkeypatch):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc.builders import pdf

    seen = {}
    real = weasyprint.HTML.render

    def spy(self, *a, **k):
        seen["document"] = real(self, *a, **k)
        return seen["document"]

    monkeypatch.setattr(weasyprint.HTML, "render", spy)
    (tmp_path / ".git").mkdir()
    captured = {}
    real_write = pdf._write_pdf
    monkeypatch.setattr(
        pdf,
        "_write_pdf",
        lambda html, base, out, options: (
            captured.update(html=html),
            real_write(html, base, out, options),
        ),
    )
    pdf.build(
        DOC,
        {"review_markup": "show", "review_author": "Greg Shallard"},
        tmp_path / "o.pdf",
        repo_root=tmp_path,
        doc_path=tmp_path / "d.md",
    )
    assert "ins.critic-ins" in captured["html"] and 'class="critic-ins"' in captured["html"]
    page = seen["document"].pages[0]._page_box
    assert not any(type(c).__name__ == "FootnoteAreaBox" for c in page.children)
    texts = " ".join(all_texts(page))
    assert "Check this" in texts and "Greg Shallard" in texts  # a margin note, with its author
    assert "critic-margin-note" in captured["html"] and "REVIEW" in texts
    assert "5 changes" in texts and "1 comment" in texts
    assert not any(type(c).__name__ == "FootnoteAreaBox" for c in page.children)

    pdf.build(DOC, {}, tmp_path / "o.pdf", repo_root=tmp_path, doc_path=tmp_path / "d.md")
    assert "critic-" not in captured["html"] and "Check this" not in captured["html"]


# ---------------------------------------------------------------------------
# Extract, lint, schema, CLI
# ---------------------------------------------------------------------------


def test_tracked_changes_round_trip_through_extract(tmp_path):
    from md_doc.extractors import extract_file

    text = extract_file(str(build_docx(tmp_path, "show")))
    assert "{--old--}" in text
    assert "{~~colour~>color~~}" in text
    assert "{==world==}{>>Check this<<}" in text
    assert "{++New bullet++}" in text


def test_extract_leaves_clean_documents_alone(tmp_path):
    from md_doc.extractors import extract_file

    text = extract_file(str(build_docx(tmp_path, "accept")))
    assert "{++" not in text and "{--" not in text and "color" in text


def lint(tmp_path: Path, frontmatter: str, body: str):
    from md_doc.linter import lint_file

    (tmp_path / ".git").mkdir(exist_ok=True)
    doc = tmp_path / "d.md"
    doc.write_text(f"---\ntitle: T\n{frontmatter}---\n{body}")
    return lint_file(doc, repo_root=tmp_path)


def test_lint_flags_broken_marks_and_final_documents(tmp_path):
    issues = lint(tmp_path, "", "# T\n\nbad {++ mark\n")
    assert any(i.severity == "error" and "never closed" in i.message for i in issues)
    issues = lint(tmp_path, "status: final\n", "# T\n\nOk {++add++}\n")
    assert [i.severity for i in issues] == ["warning"] and "final" in issues[0].message
    assert lint(tmp_path, "status: draft\n", "# T\n\nOk {++add++}\n") == []
    assert lint(tmp_path, "status: final\n", "# T\n\n`{++code++}`\n") == []


def test_schema_checks_the_review_keys():
    from md_doc.config_schema import validate_config

    assert validate_config({"review_markup": "show", "review_author": "Ada"}) == []
    assert validate_config({"review_markup": "maybe"})[0][0] == "error"


def test_cli_review_flag_overrides_the_config(tmp_path):
    (tmp_path / ".git").mkdir()
    doc = tmp_path / "d.md"
    doc.write_text("---\ntitle: T\noutputs: [docx]\n---\n# T\n\nHello {++there++}.\n")
    runner = CliRunner()
    result = runner.invoke(main, ["build", str(doc), "--review", "show", "--force"])
    assert result.exit_code == 0, result.output
    assert "<w:ins " in xml_of(tmp_path / "d.docx")
    result = runner.invoke(main, ["build", str(doc), "--force"])
    assert result.exit_code == 0, result.output
    assert "<w:ins " not in xml_of(tmp_path / "d.docx")


# ---------------------------------------------------------------------------
# simple mode and review_exports
# ---------------------------------------------------------------------------

BAR = '<span class="critic-bar"></span>'


def test_simple_applies_changes_and_marks_them():
    out = critic.apply(SAMPLE, "simple")
    assert f"# {BAR}New title" in out
    assert f"Hello {BAR}{BAR}color " in out  # deletion marker, then the new text
    assert f"- {BAR}added item\n- keep\n" in out
    assert "gone item" not in out and "removed" not in out
    assert '<mark class="critic-hl critic-c1">world</mark>' in out
    assert 'data-comment="check *this*"' in out
    assert "<ins" not in out and "<del" not in out


def test_simple_drops_wholly_deleted_lines():
    out = critic.apply("a\n\n{--\nremoved para\n--}\n\n- {--gone--}\n- keep\n", "simple")
    assert out == "a\n\n- keep\n"


def test_simple_is_accept_for_slides():
    assert critic.apply("{++a++}{--b--}", "simple", fmt="pptx") == "a"


def test_changed_blocks_get_a_class():
    html = critic.mark_changed_blocks(
        f'<h1>{BAR}T</h1><ul><li>a<ul><li>{BAR}b</li></ul></li></ul><p class="x">{BAR}q</p><p>z</p>'
    )
    assert '<h1 class="critic-changed">' in html
    assert '<li class="critic-changed">' in html and html.count("critic-changed") == 3
    assert '<p class="critic-changed x">' in html and "<p>z</p>" in html


def test_pdf_simple_draws_change_bars_and_keeps_comments(tmp_path, monkeypatch):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc.builders import pdf

    captured = {}
    real_write = pdf._write_pdf
    monkeypatch.setattr(
        pdf,
        "_write_pdf",
        lambda html, base, out, opts: (
            captured.update(html=html),
            real_write(html, base, out, opts),
        ),
    )
    (tmp_path / ".git").mkdir()
    pdf.build(
        DOC,
        {"review_markup": "simple"},
        tmp_path / "o.pdf",
        repo_root=tmp_path,
        doc_path=tmp_path / "d.md",
    )
    html = captured["html"]
    assert 'class="critic-changed"' in html and '<tr class="critic-changed">' in html
    assert "critic-margin-note" in html and "Check this" in html
    assert "Changed" in html and "critic-legend-bar" in html
    assert "<ins" not in html.split("<body>")[1] and weasyprint is not None


def test_word_simple_has_change_bars_comments_and_no_revisions(tmp_path):
    from docx import Document

    out = build_docx(tmp_path, "simple")
    xml = xml_of(out)
    assert "<w:ins " not in xml and "<w:del " not in xml
    assert 'w:right w:val="single"' in xml and 'w:highlight w:val="yellow"' in xml
    assert [c.text for c in Document(str(out)).comments] == ["Check this"]


def test_show_titles_are_clean_text(tmp_path, monkeypatch):
    from md_doc.builders import pdf

    assert pdf._extract_title('# <ins class="critic-ins">New</ins> title\n') == "New title"


def write_variants_doc(tmp_path: Path, exports: str) -> Path:
    (tmp_path / ".git").mkdir(exist_ok=True)
    doc = tmp_path / "report.md"
    doc.write_text(
        f"---\ntitle: T\noutputs: [docx]\nreview_exports: {exports}\n---\n"
        "# T\n\nHello {--old --}{~~colour~>color~~} there.\n"
    )
    return doc


def test_review_exports_writes_one_file_per_variant(tmp_path):
    doc = write_variants_doc(tmp_path, "[clean, original, simple, markup]")
    result = CliRunner().invoke(main, ["build", str(doc), "--force"])
    assert result.exit_code == 0, result.output
    names = sorted(p.name for p in tmp_path.glob("*.docx"))
    assert names == [
        "report-clean.docx",
        "report-markup.docx",
        "report-original.docx",
        "report-simple.docx",
    ]
    assert "<w:ins " in xml_of(tmp_path / "report-markup.docx")
    assert "<w:ins " not in xml_of(tmp_path / "report-clean.docx")
    assert "<w:right " in xml_of(tmp_path / "report-simple.docx")


def test_review_flag_replaces_review_exports(tmp_path):
    doc = write_variants_doc(tmp_path, "[clean, markup]")
    result = CliRunner().invoke(main, ["build", str(doc), "--review", "show", "--force"])
    assert result.exit_code == 0, result.output
    assert [p.name for p in tmp_path.glob("*.docx")] == ["report.docx"]
    assert "<w:ins " in xml_of(tmp_path / "report.docx")


def test_review_exports_are_incremental_and_validated(tmp_path):
    from md_doc.config_schema import validate_config

    doc = write_variants_doc(tmp_path, "full")
    first = CliRunner().invoke(main, ["build", str(doc)])
    again = CliRunner().invoke(main, ["build", str(doc)])
    assert "report-markup.docx" in first.output and "up to date" in again.output
    assert validate_config({"review_exports": ["clean", "simple"]}) == []
    assert validate_config({"review_exports": ["nope"]})[0][0] == "error"
    assert validate_config({"review_markup": "simple"}) == []


def test_changes_inside_a_callout_mark_the_whole_box():
    html = critic.mark_changed_blocks(
        f"<blockquote><p>a</p><ul><li>{BAR}b</li></ul></blockquote><blockquote><p>c</p></blockquote>"
    )
    assert html.startswith('<blockquote class="critic-changed">')
    assert html.count("critic-changed") == 1 and "<li>" in html


def test_word_bars_share_one_margin_strip_even_in_callouts(tmp_path):
    from md_doc.builders import docx

    (tmp_path / ".git").mkdir(exist_ok=True)
    (tmp_path / "_pdf-theme.css").write_text(
        "blockquote { background: #eef5fa; border-left: 3pt solid #2e86c1; padding: 6pt 10pt; }"
    )
    out = tmp_path / "c.docx"
    docx.build(
        "# T\n\n> **Note**\n>\n> Callout with {++new text++}.\n\nPlain {++new++} line.\n\n"
        "- {++item++}\n\n| A |\n|---|\n| {++x++} |\n",
        {"title": "T", "review_markup": "simple"},
        out,
        repo_root=tmp_path,
        doc_path=tmp_path / "d.md",
    )
    xml = xml_of(out)
    bar = '<w:right w:val="single" w:sz="18" w:space="14" w:color="CF222E"/>'
    assert xml.count(bar) == 4  # both callout paragraphs, the plain paragraph, the list item
    assert '<w:right w:val="single" w:sz="18" w:space="0" w:color="CF222E"/>' in xml  # table row
    assert '<w:left w:val="single" w:sz="12"' not in xml


def test_pdf_bars_are_drawn_in_one_margin_strip(tmp_path, monkeypatch):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc.builders import pdf

    seen = {}
    real_write = pdf._write_pdf

    def write(html, base, out, options):
        seen["options"] = options
        real_write(html, base, out, {**options, "uncompressed_pdf": True})

    monkeypatch.setattr(pdf, "_write_pdf", write)
    real_render = weasyprint.HTML.render
    monkeypatch.setattr(
        weasyprint.HTML,
        "render",
        lambda self, *a, **k: seen.setdefault("document", real_render(self, *a, **k)),
    )
    (tmp_path / ".git").mkdir()
    out = tmp_path / "o.pdf"
    pdf.build(
        "# T\n\nPlain {++new++} line.\n\n- keep\n- {++item++}\n\n"
        "> **Note** {++box++}\n\n| A | B |\n|---|---|\n| 1 | {++x++} |\n",
        {"review_markup": "simple"},
        out,
        repo_root=tmp_path,
        doc_path=tmp_path / "d.md",
    )
    data = out.read_bytes().decode("latin-1")
    rects = re.findall(r"([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) re f", data)
    bars = [r for r in rects if abs(float(r[2]) - 3.5 * 0.75) < 0.01]
    assert len(bars) == 4  # paragraph, list item, callout (whole box), table row
    assert len({r[0] for r in bars}) == 1  # one x: a single strip
    page = seen["document"].pages[0]._page_box
    content_right_pt = (page.content_box_x() + page.width) * 0.75
    assert float(bars[0][0]) > content_right_pt  # beyond the text edge, in the margin
    assert float(bars[0][0]) < seen["document"].pages[0].width * 0.75


def test_pdf_bars_follow_blocks_across_pages(tmp_path):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc import pdf_features

    html = "<style>@page{size:100mm 60mm;margin:10mm}p{margin:0 0 4mm}</style>" + "".join(
        f'<p class="critic-changed">line {i}</p>' if i in (1, 8) else f"<p>line {i}</p>"
        for i in range(12)
    )
    document = weasyprint.HTML(string=html).render()
    per_page = [pdf_features._changed_spans(p) for p in document.pages]
    assert len(document.pages) > 1 and sum(len(s) for s in per_page) == 2
    assert per_page[0] and per_page[-1]


# ---------------------------------------------------------------------------
# Visible comments and the legend
# ---------------------------------------------------------------------------


def COMMENT_SPAN(text: str, colour: int = 1, mark: int = 1) -> str:
    return (
        f'<span class="critic-comment" data-comment="{text}" data-c="{colour}" '
        f'data-mark="{mark}"></span>'
    )


class _CommentFormat:
    def format(self, text: str, colour: int = 1, mark: int = 1) -> str:
        return COMMENT_SPAN(text, colour, mark)


COMMENT = _CommentFormat()


def test_counts_ignore_code_and_frontmatter():
    found = critic.counts(SAMPLE)
    assert found == {"changes": 6, "highlights": 1, "comments": 1}
    assert critic.counts("plain `{++x++}` text") == {"changes": 0, "highlights": 0, "comments": 0}


def test_margin_notes_make_one_colour_matched_card_per_block():
    html = (
        f"<p>one {COMMENT.format('first &amp; <b>', 1)} and {COMMENT.format('again', 2, mark=0)}</p>"
        f"<p>two {COMMENT.format('second', 3)}</p>"
    )
    out = critic.margin_notes(html, "Ada <L>")
    assert "critic-comment" not in out and "critic-ref" not in out  # no numbers anywhere
    assert out.count("critic-card") == 2  # one card per paragraph, not per comment
    assert out.startswith('<p><span class="critic-margin-note critic-card" data-note="1">')
    assert out.count("critic-entry") == 3
    assert '<span class="critic-swatch critic-c1"></span><strong>Ada &lt;L&gt;</strong>' in out
    assert "first &amp; &lt;b&gt;" in out and 'critic-c2"></span><strong>' in out
    # A tick marks a comment with nothing tinted; one that follows a highlight needs none.
    assert '<span class="critic-tick critic-c1"></span>' in out
    assert (
        '<span class="critic-tick critic-c3"></span>' in out and "critic-tick critic-c2" not in out
    )


def test_table_notes_ride_in_the_first_cell_of_their_row():
    html = f"<table><tr><td>a</td><td>b {COMMENT.format('hi', 4)}</td></tr></table>"
    out = critic.margin_notes(html, "R")
    assert out.index("<td>") < out.index("critic-card") < out.index("<td>b")
    assert "<tr><span" not in out


def test_legend_describes_the_mode_and_is_absent_when_clean():
    summary = {"changes": 3, "highlights": 1, "comments": 1}
    show = critic.legend_html("show", summary, "Ada")
    assert "Ada" in show and "3 changes, 1 comment" in show and "critic-ins" in show
    simple = critic.legend_html("simple", summary, "Ada")
    assert "critic-legend-bar" in simple and "critic-ins" not in simple
    assert "critic-legend-margin" in critic.legend_html("show", summary, "Ada", margin=True)
    assert "critic-legend-margin" not in show
    assert critic.legend_html("accept", summary, "Ada") == ""
    assert critic.legend_html("show", {"changes": 0, "highlights": 0, "comments": 0}, "Ada") == ""
    assert "Review markup · Ada · 3 changes, 1 comment" in critic.legend_text(
        "show", summary, "Ada"
    )
    assert "red bar" in critic.legend_text("simple", summary, "Ada")


def test_word_gets_a_legend_only_for_marked_up_versions(tmp_path):
    from docx import Document

    show = Document(str(build_docx(tmp_path, "show")))
    assert show.paragraphs[0].text.startswith("Review markup · Greg Shallard · ")
    assert "accept or reject them on Word's Review tab" in show.paragraphs[0].text
    simple = Document(str(build_docx(tmp_path, "simple")))
    assert "red bar in the right margin" in simple.paragraphs[0].text
    clean = Document(str(build_docx(tmp_path, "accept")))
    assert not clean.paragraphs[0].text.startswith("Review markup")


def test_pdf_clean_versions_have_no_legend_or_notes(tmp_path, monkeypatch):
    from md_doc.builders import pdf

    captured = {}
    real_write = pdf._write_pdf
    monkeypatch.setattr(
        pdf,
        "_write_pdf",
        lambda html, base, out, opts: (
            captured.update(html=html),
            real_write(html, base, out, opts),
        ),
    )
    (tmp_path / ".git").mkdir(exist_ok=True)
    for mode in ("accept", "reject"):
        pdf.build(
            DOC,
            {"review_markup": mode},
            tmp_path / "o.pdf",
            repo_root=tmp_path,
            doc_path=tmp_path / "d.md",
        )
        assert "critic-legend" not in captured["html"] and "critic-margin" not in captured["html"]


# ---------------------------------------------------------------------------
# The review column: shrunk page, margin notes, no change to the document layout
# ---------------------------------------------------------------------------


def test_review_page_keeps_the_text_area_and_scales_back_to_paper_width():
    geometry = critic.review_page((210.0, 297.0), (25.0, 20.0), 20.0)
    width, height = geometry["size"]
    assert width == 210.0 + critic.COLUMN_MM
    assert geometry["margin_right"] == 20.0 + critic.COLUMN_MM
    assert round(width * geometry["zoom"], 6) == 210.0  # the sheet is still the paper width
    assert round(height * geometry["zoom"], 6) == 297.0  # ...and the paper height
    # Text area unchanged: page height - bottom margin = original page height - 20mm.
    assert round(height - geometry["margin_bottom"], 6) == 277.0
    css = critic.review_page_css(geometry)
    assert "@page { size: 282.000mm" in css and "position: absolute" in css


def test_notes_beside_the_same_line_are_spread_apart():
    weasyprint = pytest.importorskip("weasyprint")
    note = '<span class="critic-margin-note n" data-note="{0}">note {0}</span>'
    html = (
        "<style>@page{size:200mm 100mm;margin:10mm 80mm 10mm 20mm}"
        ".n{position:absolute;left:100mm;width:60mm;font-size:8pt}</style>"
        f"<p>text {note.format(1)} more {note.format(2)} and {note.format(3)}</p>"
    )
    spread = critic.spread_margin_notes(
        html, lambda markup: weasyprint.HTML(string=markup).render()
    )
    # All three wanted the same line: they are spread around it, not piled below it.
    moved = re.findall(r'data-note="(\d)" style="margin-top: (-?[\d.]+)px"', spread)
    assert len(moved) >= 2 and min(float(d) for _, d in moved) < 0 < max(float(d) for _, d in moved)
    boxes = []
    document = weasyprint.HTML(string=spread).render()

    def visit(box):
        element = getattr(box, "element", None)
        if element is not None and "critic-margin-note" in (element.get("class") or ""):
            boxes.append((box.position_y + box.margin_top, box.border_height()))
            return
        for child in getattr(box, "children", []) or []:
            visit(child)

    visit(document.pages[0]._page_box)
    boxes.sort()
    assert len(boxes) == 3
    assert all(a[0] + a[1] <= b[0] + 0.5 for a, b in zip(boxes, boxes[1:]))


def test_spread_leaves_documents_without_notes_alone():
    assert (
        critic.spread_margin_notes("<p>plain</p>", lambda _: pytest.fail("measured"))
        == "<p>plain</p>"
    )


def test_pdf_with_comments_is_paper_size_and_paginates_like_the_clean_version(tmp_path):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc.builders import pdf

    (tmp_path / ".git").mkdir(exist_ok=True)
    body = "# Review\n\n" + "".join(
        f"Paragraph {i} with {{++added words ++}}and a {{==point==}}{{>>comment {i}<<}} to check.\n\n"
        for i in range(60)
    )
    layouts = {}
    real_render = weasyprint.HTML.render
    for mode in ("accept", "simple"):
        captured = {}
        weasyprint.HTML.render = lambda self, *a, **k: captured.setdefault(
            "document", real_render(self, *a, **k)
        )
        try:
            pdf.build(
                body,
                {"review_markup": mode, "pdf_uncompressed": True},
                tmp_path / f"{mode}.pdf",
                repo_root=tmp_path,
                doc_path=tmp_path / "d.md",
            )
        finally:
            weasyprint.HTML.render = real_render
        document = captured["document"]
        ys = []
        for page in document.pages:
            for box in page._page_box.descendants():
                element = getattr(box, "element", None)
                if (
                    element is not None
                    and element.tag == "p"
                    and type(box).__name__ == "BlockBox"
                    and "critic-margin-note" not in (element.get("class") or "")
                ):
                    ys.append(round(box.position_y, 1))
        layouts[mode] = (len(document.pages), sorted(set(ys)))
    assert layouts["accept"][0] > 1 and layouts["accept"] == layouts["simple"]
    text = (tmp_path / "simple.pdf").read_bytes()
    assert re.search(rb"/MediaBox \[[^\]]*595\.27", text)  # A4 width, scaled back


# ---------------------------------------------------------------------------
# Hybrid: light highlight + bars in the text, the full changes in the review column
# ---------------------------------------------------------------------------


def test_hybrid_applies_changes_with_a_light_highlight_and_a_note_anchor():
    out = critic.apply(SAMPLE, "hybrid")
    assert re.search(rf'# {BAR}<span class="critic-chg critic-c1">New</span>', out)
    assert 'data-kind="sub" data-old="colour" data-new="color"' in out
    assert 'data-kind="del" data-old="old"' in out
    assert '<mark class="critic-hl critic-c' in out and 'data-comment="check *this*"' in out
    assert "<ins" not in out and "<del" not in out and "old " not in out.split("# ")[1]


def test_hybrid_keeps_a_deleted_line_as_a_note_on_the_line_above():
    out = critic.apply("a para\n\n- keep\n- {--gone item--}\n\nnext\n", "hybrid")
    assert "gone item</" not in out and "- keep<span" in out
    assert 'data-kind="del" data-old="gone item"' in out
    assert out.count("\n- ") == 1  # no empty bullet left behind
    table = critic.apply("| a | b |\n|---|---|\n| 1 | 2 |\n{--\n| 3 | 4 |\n--}\n", "hybrid")
    assert "|" in table and table.rstrip().endswith("|")
    assert 'data-kind="del"' in table


def test_hybrid_margin_notes_spell_out_each_change_in_distinct_colours():
    html = critic.apply("A {~~old~>new~~} B {--gone--} C {++more++} {>>why<<}", "hybrid")
    out = critic.margin_notes(f"<p>{html}</p>", "Ada")
    assert out.count("critic-card") == 1 and out.count("critic-entry") == 4
    assert [int(c) for c in re.findall(r'<span class="critic-swatch critic-c(\d+)', out)] == [
        1,
        2,
        3,
        4,
    ]
    assert '<del class="critic-del">old</del> <ins class="critic-ins">new</ins>' in out
    assert (
        "CHANGED" in out
        and "DELETED" in out
        and "ADDED" in out
        and "<strong>Ada</strong> why" in out
    )
    assert 'class="critic-chg critic-c1">new</span>' in out  # the text carries the same colour
    assert 'class="critic-chg critic-c3">more</span>' in out
    assert out.count("critic-tick") == 2  # the deletion and the stand-alone comment
    assert "critic-change-note" not in out and "critic-comment" not in out


def test_colours_cycle_and_a_highlight_shares_its_comments_colour():
    marks = "".join(f"{{++a{i}++}}" for i in range(10)) + "{==x==}{>>c<<}"
    out = critic.apply(marks, "hybrid")
    assert 'critic-chg critic-c10">a9' in out  # ten colours, then they repeat
    assert 'mark class="critic-hl critic-c1">x' in out
    assert 'data-comment="c" data-c="1" data-mark="0"' in out
    # A highlight with no comment stays the plain yellow.
    assert '<mark class="critic-hl">y</mark>' in critic.apply("{==y==}", "hybrid")


def test_hybrid_legend_and_modes_per_format():
    summary = {"changes": 2, "highlights": 0, "comments": 0}
    assert "matched to its margin note" in critic.legend_html("hybrid", summary, "Ada")
    assert critic.mode_for({"review_markup": "hybrid"}, "pdf") == "hybrid"
    assert critic.mode_for({"review_markup": "hybrid"}, "docx") == "show"
    assert critic.mode_for({"review_markup": "hybrid"}, "pptx") == "accept"
    assert critic.apply("{++a++}{--b--}", "hybrid", fmt="pptx") == "a"
    assert critic.VARIANTS["hybrid"] == ("hybrid", "-hybrid")


def test_pdf_hybrid_has_a_review_column_even_without_comments(tmp_path, monkeypatch):
    weasyprint = pytest.importorskip("weasyprint")
    from md_doc.builders import pdf

    seen = {}
    real_render = weasyprint.HTML.render
    monkeypatch.setattr(
        weasyprint.HTML,
        "render",
        lambda self, *a, **k: seen.update(doc=real_render(self, *a, **k)) or seen["doc"],
    )
    (tmp_path / ".git").mkdir(exist_ok=True)
    pdf.build(
        "# T\n\nWe take {~~six~>eight~~} weeks and {--roughly --}cost $5.\n",
        {"review_markup": "hybrid", "review_author": "Ada"},
        tmp_path / "o.pdf",
        repo_root=tmp_path,
        doc_path=tmp_path / "d.md",
    )
    texts = " ".join(all_texts(seen["doc"].pages[0]._page_box))
    assert "CHANGED" in texts and "DELETED" in texts and "REVIEW" in texts
    assert seen["doc"].pages[0].width > 210 * 96 / 25.4  # the page was widened for the column


def test_word_hybrid_is_tracked_changes(tmp_path):
    xml = xml_of(build_docx(tmp_path, "hybrid"))
    assert "<w:ins " in xml and "<w:del " in xml


def test_schema_and_variants_accept_hybrid(tmp_path):
    from md_doc.config_schema import validate_config

    assert validate_config({"review_markup": "hybrid", "review_exports": ["hybrid", "clean"]}) == []
    doc = write_variants_doc(tmp_path, "[hybrid]")
    result = CliRunner().invoke(main, ["build", str(doc), "--force"])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "report-hybrid.docx").exists()
