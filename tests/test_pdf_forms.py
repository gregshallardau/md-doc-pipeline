"""Tests for the PDF forms build-out.

Covers the ?[...] shorthand (yesno, box grids, attrs), the REAL AcroForm
output (integration — not mocked), the finisher that patches metadata
WeasyPrint drops, form linting, and the dotx Word-form-field mapping.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from md_doc.builders.pdf import _expand_form_fields, _field_to_html
from md_doc.forms import iter_field_specs, parse_field_spec


@pytest.fixture()
def tmp_repo(tmp_path):
    (tmp_path / ".git").mkdir()
    return tmp_path


# ── shorthand → HTML ─────────────────────────────────────────────────────────


class TestFieldToHtml:
    def test_yesno_pair(self):
        html = _field_to_html("yesno: cover")
        assert 'name="cover_yes"' in html and 'name="cover_no"' in html
        assert html.count('type="checkbox"') == 2

    def test_checkbox_checked_and_value(self):
        html = _field_to_html("checkbox: agree, label=I agree, checked, value=YES")
        assert "checked" in html and 'value="YES"' in html and "I agree" in html

    def test_tel_and_url_types(self):
        assert 'type="tel"' in _field_to_html("tel: phone")
        assert 'type="url"' in _field_to_html("url: website")

    def test_textarea_extra_attrs(self):
        html = _field_to_html("textarea: notes, rows=3, title=Hint here")
        assert 'rows="3"' in html and 'title="Hint here"' in html

    def test_unsafe_attr_name_dropped(self):
        html = _field_to_html("text: x, onfocus=alert(1)")
        assert "onfocus" not in html


class TestBoxConstruct:
    def test_box_renders_field_grid(self):
        md = "?[box]\n**Name** ?[text: name]\n**City** ?[text: city] | **State** ?[text: state]\n?[/box]\n"
        html = _expand_form_fields(md, is_form=True)
        assert 'class="field-box"' in html
        assert "<strong>Name</strong>" in html  # markdown label converted
        assert 'name="city"' in html and 'name="state"' in html

    def test_box_short_row_gets_colspan(self):
        md = "?[box]\nOne | Two | Three\nFull width row\n?[/box]\n"
        html = _expand_form_fields(md, is_form=True)
        assert 'colspan="3"' in html

    def test_box_widths(self):
        md = "?[box: widths=72,28]\nQuestion | ?[yesno: q1]\n?[/box]\n"
        html = _expand_form_fields(md, is_form=True)
        assert "width: 72.00%" in html and "width: 28.00%" in html

    def test_box_italic_hint(self):
        md = "?[box]\nLabel *small hint* ?[text: x]\n?[/box]\n"
        html = _expand_form_fields(md, is_form=True)
        assert "<em>small hint</em>" in html

    def test_row_labels_converted(self):
        md = "?[row]\n**City** ?[text: city] | **Post** ?[text: post]\n?[/row]\n"
        html = _expand_form_fields(md, is_form=True)
        assert "<strong>City</strong>" in html

    def test_unclosed_box_does_not_hang(self):
        md = "?[box]\nno close marker\n"
        html = _expand_form_fields(md, is_form=True)
        assert isinstance(html, str)


# ── forms.py spec parsing ────────────────────────────────────────────────────


class TestSpecParsing:
    def test_parse_basic(self):
        assert parse_field_spec("text: name, required")[0:2] == ("text", "name")

    def test_parse_select_options(self):
        ftype, name, options, _ = parse_field_spec("select: region | North | South")
        assert (ftype, name, options) == ("select", "region", ["North", "South"])

    def test_structural_markers_skipped(self):
        assert parse_field_spec("row") is None
        assert parse_field_spec("/box") is None
        assert parse_field_spec("box: widths=70,30") is None
        assert parse_field_spec("submit Send") is None

    @pytest.mark.parametrize(
        "label",
        [
            "$10,000,000",
            '"$10,000,000"',
            r"$10\,000\,000",
            "'Cover, including tax'",
            r'"Say \"yes\", then continue"',
            "Customer's cover",
        ],
    )
    def test_delimited_labels_preserve_following_flags(self, label):
        from md_doc.forms import parse_field_attrs

        expected = {
            "$10,000,000": "$10,000,000",
            '"$10,000,000"': "$10,000,000",
            r"$10\,000\,000": "$10,000,000",
            "'Cover, including tax'": "Cover, including tax",
            r'"Say \"yes\", then continue"': 'Say "yes", then continue',
            "Customer's cover": "Customer's cover",
        }[label]
        spec = f"checkbox: amount, label={label}, required, checked, value=YES"
        parsed = parse_field_spec(spec)
        assert parsed[3] == {"label": expected, "required": True, "checked": True, "value": "YES"}
        assert parse_field_attrs(r"title=C:\Temp\cover, value=10,000") == {
            "title": r"C:\Temp\cover",
            "value": "10,000",
        }
        html = _field_to_html(spec)
        from html import unescape

        assert expected in unescape(html)
        assert 'required checked value="YES"' in html

    def test_iter_skips_structure(self):
        md = "?[box]\n?[text: a] | ?[yesno: b]\n?[/box]\n?[submit Go]\n"
        specs = iter_field_specs(md)
        assert [(s[0], s[1]) for s in specs] == [("text", "a"), ("yesno", "b")]


# ── REAL AcroForm output (integration, not mocked) ───────────────────────────


class TestAcroFormIntegration:
    def _build(self, tmp_repo: Path, body: str) -> Path:
        pytest.importorskip("weasyprint")
        from md_doc.builders.pdf import build

        doc = tmp_repo / "form.md"
        md = f"---\ntitle: F\n---\n\n# F\n\n{body}"
        doc.write_text(md, encoding="utf-8")
        out = tmp_repo / "form.pdf"
        build(md, {"title": "F", "pdf_forms": True}, out, doc_path=doc, repo_root=tmp_repo)
        return out

    def test_currency_label_is_visible_in_pdf(self, tmp_repo):
        import pdfplumber

        out = self._build(tmp_repo, 'Amount: ?[checkbox: amount, label="$10,000,000", required]')
        with pdfplumber.open(out) as pdf:
            assert "$10,000,000" in "\n".join(page.extract_text() or "" for page in pdf.pages)

    def test_fields_are_real_acroform(self, tmp_repo):
        pdfium = pytest.importorskip("pypdfium2")
        import pypdfium2.raw as raw

        out = self._build(
            tmp_repo,
            "**Name** ?[text: name, required, title=Full legal name]\n\n"
            "**Locked** ?[text: locked, readonly]\n\n"
            "?[yesno: cover]\n",
        )
        doc = pdfium.PdfDocument(str(out))
        assert raw.FPDF_GetFormType(doc.raw) == 1  # real AcroForm
        doc.init_forms()
        form = doc.formenv
        flags = {}
        widgets = 0
        for pi in range(len(doc)):
            page = doc[pi]
            for ai in range(raw.FPDFPage_GetAnnotCount(page.raw)):
                a = raw.FPDFPage_GetAnnot(page.raw, ai)
                if raw.FPDFAnnot_GetSubtype(a) == raw.FPDF_ANNOT_WIDGET:
                    widgets += 1
                    import ctypes

                    fn = raw.FPDFAnnot_GetFormFieldName
                    m = fn(form.raw, a, None, 0)
                    buf = ctypes.create_string_buffer(max(m, 2))
                    fn(form.raw, a, ctypes.cast(buf, ctypes.POINTER(raw.FPDF_WCHAR)), m)
                    name = buf.raw[: max(m - 2, 0)].decode("utf-16-le", errors="replace")
                    flags[name] = raw.FPDFAnnot_GetFormFieldFlags(form.raw, a)
                raw.FPDFPage_CloseAnnot(a)
        assert widgets == 4  # name, locked, cover_yes, cover_no
        # finisher-patched metadata: required (bit 2) and readonly (bit 1)
        assert flags["name"] & 2, "required flag missing"
        assert flags["locked"] & 1, "readonly flag missing"

    def test_tooltip_reaches_pdf(self, tmp_repo):
        # Exercise the collect+finisher machinery directly with an
        # uncompressed write so /TU (tooltip) and /Ff are byte-visible.
        weasyprint = pytest.importorskip("weasyprint")
        from md_doc.builders.pdf import _collect_form_field_meta, _make_forms_finisher

        html = (
            "<html><head><style>input{appearance:auto;width:100pt;height:12pt}"
            "</style></head><body><form>"
            '<input type="text" name="who" title="Full legal name" required>'
            "</form></body></html>"
        )
        meta = _collect_form_field_meta(html)
        assert meta["who"] == {"required": True, "tooltip": "Full legal name"}
        pdf = weasyprint.HTML(string=html).write_pdf(
            pdf_forms=True, uncompressed_pdf=True, finisher=_make_forms_finisher(meta)
        )
        assert b"/TU (Full legal name)" in pdf
        assert b"/Ff 2" in pdf  # required bit

    def test_form_support_css_only_for_forms(self, tmp_repo):
        pytest.importorskip("weasyprint")
        from unittest.mock import MagicMock, patch

        from md_doc.builders.pdf import build

        doc = tmp_repo / "d.md"
        doc.write_text("# T\n\nplain\n", encoding="utf-8")
        with patch("md_doc.builders.pdf.weasyprint") as wp:
            wp.HTML.return_value = MagicMock()
            build("# T\n\nplain\n", {"title": "T"}, tmp_repo / "d.pdf", doc_path=doc)
            html = wp.HTML.call_args.kwargs["string"]
        assert "field-box" not in html
        # …but the first-H1 fix is always present
        assert "h1:first-of-type" in html


# ── linting ──────────────────────────────────────────────────────────────────


class TestFormLint:
    def _lint(self, tmp_repo: Path, md: str):
        from md_doc.linter import lint_file

        doc = tmp_repo / "f.md"
        doc.write_text(md, encoding="utf-8")
        return [i.message for i in lint_file(doc, repo_root=tmp_repo)]

    def test_unknown_type_warns(self, tmp_repo):
        msgs = self._lint(tmp_repo, "---\ntitle: T\npdf_forms: true\n---\n\n?[texd: x]\n")
        assert any("Unknown form field type 'texd'" in m for m in msgs)

    def test_duplicate_names_warn(self, tmp_repo):
        msgs = self._lint(
            tmp_repo, "---\ntitle: T\npdf_forms: true\n---\n\n?[text: a]\n\n?[text: a]\n"
        )
        assert any("Duplicate form field name 'a'" in m for m in msgs)

    def test_missing_pdf_forms_warns(self, tmp_repo):
        msgs = self._lint(tmp_repo, "---\ntitle: T\n---\n\n?[text: a]\n")
        assert any("pdf_forms: true" in m for m in msgs)

    def test_clean_form_no_warnings(self, tmp_repo):
        msgs = self._lint(
            tmp_repo,
            "---\ntitle: T\npdf_forms: true\n---\n\n?[text: a]\n\n?[yesno: b]\n",
        )
        assert not any("form field" in m.lower() for m in msgs)


# ── dotx mapping ─────────────────────────────────────────────────────────────


class TestDotxFormFields:
    def _build_dotx(self, tmp_repo: Path, body: str, output_format: str = "dotx") -> str:
        from md_doc.builders.docx import build

        doc = tmp_repo / "t.md"
        md = f"---\ntitle: T\n---\n\n# T\n\n{body}"
        doc.write_text(md, encoding="utf-8")
        out = tmp_repo / f"t.{output_format}"
        build(
            md,
            {"title": "T", "cover_page": False},
            out,
            output_format=output_format,
            doc_path=doc,
            repo_root=tmp_repo,
        )
        with zipfile.ZipFile(out) as z:
            return z.read("word/document.xml").decode("utf-8")

    def test_text_field_maps_to_formtext(self, tmp_repo):
        xml = self._build_dotx(tmp_repo, "**Name** ?[text: full_name]\n")
        assert "FORMTEXT" in xml and 'w:val="full_name"' in xml

    def test_checkbox_maps_to_formcheckbox(self, tmp_repo):
        xml = self._build_dotx(tmp_repo, "?[checkbox: agree, label=I agree]\n")
        assert "FORMCHECKBOX" in xml and "I agree" in xml

    @pytest.mark.parametrize("output_format", ["docx", "dotx"])
    def test_currency_and_comma_labels_survive_word_export(self, tmp_repo, output_format):
        xml = self._build_dotx(
            tmp_repo, 'Amount: ?[checkbox: amount, label="$10,000,000", checked]\n', output_format
        )
        assert "$10,000,000" in xml
        if output_format == "dotx":
            assert "FORMCHECKBOX" in xml

    def test_select_maps_to_dropdown_with_options(self, tmp_repo):
        xml = self._build_dotx(tmp_repo, "?[select: region | North | South]\n")
        assert "FORMDROPDOWN" in xml
        assert 'w:val="North"' in xml and 'w:val="South"' in xml

    def test_yesno_maps_to_checkbox_pair(self, tmp_repo):
        xml = self._build_dotx(tmp_repo, "?[yesno: cover]\n")
        assert xml.count("FORMCHECKBOX") == 2
        assert 'w:val="cover_yes"' in xml and 'w:val="cover_no"' in xml

    def test_no_marker_leaks(self, tmp_repo):
        xml = self._build_dotx(tmp_repo, "?[box]\nA | ?[text: a]\n?[/box]\n\n?[submit Send]\n")
        assert "?[" not in xml

    def test_docx_renders_field_as_bordered_box_not_form_field(self, tmp_repo):
        from md_doc.builders.docx import build

        doc = tmp_repo / "t.md"
        md = "---\ntitle: T\n---\n\n# T\n\n**Name** ?[text: name]\n"
        doc.write_text(md, encoding="utf-8")
        out = tmp_repo / "t.docx"
        build(
            md,
            {"title": "T", "cover_page": False},
            out,
            output_format="docx",
            doc_path=doc,
            repo_root=tmp_repo,
        )
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        # A plain .docx shows the input as a PDF-sized bordered box, never a
        # live Word form field (those belong to the .dotx template).
        assert 'w:fill="FAFAFA"' in xml and "<w:pBdr>" in xml and "FORMTEXT" not in xml
        assert "________" not in xml


# ── first-H1 page-break fix (docx side) ──────────────────────────────────────


def test_docx_first_h1_after_letterhead_no_break(tmp_repo):
    from docx import Document

    from md_doc.builders.docx import build

    (tmp_repo / "_pdf-theme.css").write_text(
        "@page { size: A4; margin: 25mm 20mm 22mm 25mm; }\n"
        ".report-body h1 { page-break-before: always; }\n",
        encoding="utf-8",
    )
    md = "---\ntitle: T\n---\n\nLetterhead line\n\n# First Heading\n\nBody\n\n# Second Heading\n\nMore\n"
    doc_path = tmp_repo / "d.md"
    doc_path.write_text(md, encoding="utf-8")
    out = tmp_repo / "d.docx"
    build(
        md,
        {"title": "T", "cover_page": False},
        out,
        output_format="docx",
        doc_path=doc_path,
        repo_root=tmp_repo,
    )
    d = Document(str(out))
    first = next(p for p in d.paragraphs if p.text == "First Heading")
    second = next(p for p in d.paragraphs if p.text == "Second Heading")
    from tests.test_docx_parity import starts_new_page

    assert not starts_new_page(first)  # letterhead must not force page 2
    assert starts_new_page(second)  # later H1s still break


# ── adjacent tables ──────────────────────────────────────────────────────────


def test_adjacent_tables_get_separation_css(tmp_repo):
    # Theme-independent: two separate tables must never render flush (a theme
    # with no `table { margin }` made them look like one merged grid).
    pytest.importorskip("weasyprint")
    from unittest.mock import MagicMock, patch

    from md_doc.builders.pdf import build

    doc = tmp_repo / "d.md"
    doc.write_text("# T\n", encoding="utf-8")
    with patch("md_doc.builders.pdf.weasyprint") as wp:
        wp.HTML.return_value = MagicMock()
        build("# T\n", {"title": "T"}, tmp_repo / "d.pdf", doc_path=doc)
        html = wp.HTML.call_args.kwargs["string"]
    assert "table + table { margin-top" in html


def test_lint_warns_on_jammed_tables(tmp_repo):
    from md_doc.linter import lint_file

    doc = tmp_repo / "jam.md"
    doc.write_text(
        "---\ntitle: T\n---\n\n| A |\n|---|\n| 1 |\n| B |\n|---|\n| 2 |\n",
        encoding="utf-8",
    )
    msgs = [i.message for i in lint_file(doc, repo_root=tmp_repo)]
    assert any("Two tables jammed together" in m for m in msgs)


def test_lint_ok_on_separated_tables(tmp_repo):
    from md_doc.linter import lint_file

    doc = tmp_repo / "ok.md"
    doc.write_text(
        "---\ntitle: T\n---\n\n| A |\n|---|\n| 1 |\n\n| B |\n|---|\n| 2 |\n",
        encoding="utf-8",
    )
    msgs = [i.message for i in lint_file(doc, repo_root=tmp_repo)]
    assert not any("Two tables jammed together" in m for m in msgs)


@pytest.mark.parametrize("spec", ["submit Send", "submit: Send"])
def test_submit_alias_is_a_button_not_a_named_text_field(spec):
    from md_doc.forms import parse_field_spec
    from md_doc.builders.pdf import _field_to_html

    assert parse_field_spec(spec) is None
    assert _field_to_html(spec) == '<input type="submit" value="Send">'


def test_submit_colon_does_not_create_a_word_field(tmp_repo):
    from md_doc.builders.docx import _convert_form_fields_for_dotx

    assert _convert_form_fields_for_dotx("?[submit: Send]") == ""


def test_signature_row_labels_become_captions():
    md = "?[row]\n?[signature: sig] | **Date** ?[date: d]\n?[/row]\n"
    html = _expand_form_fields(md, is_form=True)
    assert '<div class="signature-label">Date</div>' in html


def test_box_question_rows_centre_and_currency_prefix_inline():
    md = "?[box]\nTotal turnover: | $ ?[number: t]\n**City** ?[text: c]\n?[/box]\n"
    html = _expand_form_fields(md, is_form=True)
    assert '<tr class="qa">' in html
    assert 'class="prefixed-field"' in html
    assert html.count('<tr class="qa">') == 1  # the labelled row stays top-aligned


def test_trailing_checkbox_text_folds_into_label():
    html = _expand_form_fields("?[checkbox: agree] I agree to terms\n", is_form=True)
    assert "option-solo" in html
    assert "I agree to terms</label>" in html


def test_option_controls_are_centred_on_label_caps():
    from md_doc.builders.pdf import _form_support_css

    css = _form_support_css(None)
    assert "position: relative; top: -0.1em" in css
